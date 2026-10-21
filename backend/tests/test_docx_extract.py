import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

import app
from app import convert_docx_to_text, read_docx


class DocxExtractTest(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmpdir, ignore_errors=True)

    def _path(self, name):
        return str(Path(self.tmpdir) / name)

    def test_table_text_kept_in_document_order(self):
        doc = Document()
        doc.add_paragraph("表格前正文")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "姓名"
        table.cell(0, 1).text = "年龄"
        table.cell(1, 0).text = "张三"
        table.cell(1, 1).text = "三十"
        doc.add_paragraph("表格后正文")
        doc.save(self._path("table.docx"))

        self.assertEqual(
            read_docx(self._path("table.docx")),
            "表格前正文\n姓名 年龄\n张三 三十\n表格后正文",
        )

    def test_merged_cells_are_not_duplicated(self):
        doc = Document()
        doc.add_paragraph("开头")
        table = doc.add_table(rows=3, cols=3)
        table.cell(0, 0).text = "合并表头"
        table.cell(0, 0).merge(table.cell(0, 2))
        table.cell(1, 0).text = "竖合并"
        table.cell(1, 0).merge(table.cell(2, 0))
        table.cell(1, 1).text = "B1"
        table.cell(1, 2).text = "C1"
        table.cell(2, 1).text = "B2"
        table.cell(2, 2).text = "C2"
        inner = table.cell(2, 2).add_table(rows=2, cols=1)
        inner.cell(0, 0).text = "嵌套甲"
        inner.cell(1, 0).text = "嵌套乙"
        doc.add_paragraph("结尾")
        doc.save(self._path("merge.docx"))

        self.assertEqual(
            read_docx(self._path("merge.docx")),
            "开头\n合并表头\n竖合并 B1 C1\nB2 C2\n嵌套甲\n嵌套乙\n结尾",
        )

    def test_inline_content_control_and_field_text_kept(self):
        doc = Document()
        paragraph = doc.add_paragraph("前")
        paragraph._p.append(
            parse_xml(
                '<w:sdt %s><w:sdtPr/><w:sdtContent><w:r><w:t>控件</w:t></w:r>'
                "</w:sdtContent></w:sdt>" % nsdecls("w")
            )
        )
        paragraph._p.append(parse_xml("<w:r %s><w:t>后</w:t></w:r>" % nsdecls("w")))
        # 删除的修订不该出现在结果里：它的文字在 w:delText 而不是 w:t
        paragraph._p.append(
            parse_xml("<w:del %s><w:r><w:delText>删掉的</w:delText></w:r></w:del>" % nsdecls("w"))
        )
        field = doc.add_paragraph("页码 ")
        field._p.append(
            parse_xml(
                '<w:fldSimple %s w:instr=" PAGE "><w:r><w:t>3</w:t></w:r></w:fldSimple>'
                % nsdecls("w")
            )
        )
        doc.save(self._path("sdt.docx"))

        self.assertEqual(read_docx(self._path("sdt.docx")), "前控件后\n页码 3")

    def test_block_content_control_contents_are_read(self):
        doc = Document()
        doc.add_paragraph("控件外")
        body = doc.element.body
        # 块级内容控件把段落和表格整个包在 w:sdtContent 里
        body.insert(
            len(body) - 1,
            parse_xml(
                "<w:sdt %s><w:sdtPr/><w:sdtContent>"
                "<w:p><w:r><w:t>控件里的正文</w:t></w:r></w:p>"
                "<w:tbl><w:tr><w:tc><w:p><w:r><w:t>控件里的表格</w:t></w:r></w:p>"
                "</w:tc></w:tr></w:tbl>"
                "</w:sdtContent></w:sdt>" % nsdecls("w")
            ),
        )
        doc.save(self._path("block_sdt.docx"))

        self.assertEqual(
            read_docx(self._path("block_sdt.docx")),
            "控件外\n控件里的正文\n控件里的表格",
        )

    def test_line_break_in_cell_is_aligned_with_neighbour_cells(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        paragraph = table.cell(0, 0).paragraphs[0]
        paragraph.add_run("A")
        paragraph.add_run().add_break()
        paragraph.add_run("B")
        table.cell(0, 1).text = "X"
        doc.save(self._path("cell_break.docx"))

        # 同一个单元格里的换行也要参与按行对齐，不能把 "A\nB" 整块当成一行
        self.assertEqual(read_docx(self._path("cell_break.docx")), "A X\nB")

    def test_internal_blank_line_in_cell_keeps_alignment(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        paragraph = table.cell(0, 0).paragraphs[0]
        paragraph.add_run("A")
        paragraph.add_run().add_break()
        paragraph.add_run().add_break()
        paragraph.add_run("B")
        table.cell(0, 1).text = "X\nY\nZ"
        doc.save(self._path("cell_blank_line.docx"))

        # 连续换行留下的空行是列对齐的占位，丢了 B 就会错配到 Y 那一行
        self.assertEqual(
            read_docx(self._path("cell_blank_line.docx")),
            "A X\nY\nB Z",
        )

    def test_content_controls_around_rows_and_cells_are_read(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "普通甲"
        table.cell(0, 1).text = "普通乙"
        table.cell(1, 0).text = "控件甲"
        table.cell(1, 1).text = "控件乙"

        tbl = table._tbl
        # 第二行整行包进内容控件（w:sdtRow）
        row = tbl.tr_lst[-1]
        tbl.remove(row)
        row_wrapper = parse_xml("<w:sdt %s><w:sdtPr/><w:sdtContent/></w:sdt>" % nsdecls("w"))
        row_wrapper.find(qn("w:sdtContent")).append(row)
        tbl.append(row_wrapper)

        # 第一行第二格单独包进内容控件（w:sdtCell）
        first_row = tbl.tr_lst[0]
        cell = first_row.tc_lst[-1]
        first_row.remove(cell)
        cell_wrapper = parse_xml("<w:sdt %s><w:sdtPr/><w:sdtContent/></w:sdt>" % nsdecls("w"))
        cell_wrapper.find(qn("w:sdtContent")).append(cell)
        first_row.append(cell_wrapper)
        doc.save(self._path("row_cell_sdt.docx"))

        self.assertEqual(
            read_docx(self._path("row_cell_sdt.docx")),
            "普通甲 普通乙\n控件甲 控件乙",
        )

    def test_hyperlink_text_is_not_dropped(self):
        doc = Document()
        paragraph = doc.add_paragraph("点击 ")
        paragraph._p.append(
            parse_xml(
                '<w:hyperlink %s r:id="rId999"><w:r><w:t>这个链接</w:t></w:r></w:hyperlink>'
                % nsdecls("w", "r")
            )
        )
        doc.save(self._path("link.docx"))

        self.assertEqual(read_docx(self._path("link.docx")), "点击 这个链接")

    def test_document_with_table_skips_pandoc(self):
        # pandoc 的 plain 输出会把表格画成一行行横线，这些横线会被当成分页标记
        doc = Document()
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "单元格"
        doc.save(self._path("table_only.docx"))

        with patch.object(app, "_pandoc_available", True), patch.object(
            app.pypandoc, "convert_file", side_effect=AssertionError("带表格时不应调用 pandoc")
        ):
            self.assertEqual(convert_docx_to_text(self._path("table_only.docx")), "单元格")

    def test_document_without_table_still_uses_pandoc(self):
        doc = Document()
        doc.add_paragraph("正文")
        doc.save(self._path("plain.docx"))

        with patch.object(app, "_pandoc_available", True), patch.object(
            app.pypandoc, "convert_file", return_value="pandoc 结果"
        ) as convert:
            self.assertEqual(convert_docx_to_text(self._path("plain.docx")), "pandoc 结果")
        convert.assert_called_once()


if __name__ == "__main__":
    unittest.main()
